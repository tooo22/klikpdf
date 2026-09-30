import fitz
from pathlib import Path
from pdf2docx import Converter
import openpyxl
import pdfplumber

class ConvertService:
    @staticmethod
    def pdf_to_word(pdf_path: Path, docx_output_path: Path) -> Path:
        cv = Converter(str(pdf_path))
        cv.convert(str(docx_output_path), start=0, end=None)
        cv.close()
        return docx_output_path

    @staticmethod
    def pdf_to_excel(pdf_path: Path, xlsx_output_path: Path) -> Path:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "KlikPDF Data"
        
        row_idx = 1
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            for col_idx, val in enumerate(row, start=1):
                                ws.cell(row=row_idx, column=col_idx, value=val or "")
                            row_idx += 1
                        row_idx += 1
                else:
                    text = page.extract_text() or ""
                    for line in text.split("\n"):
                        ws.cell(row=row_idx, column=1, value=line)
                        row_idx += 1
        wb.save(xlsx_output_path)
        return xlsx_output_path

convert_service = ConvertService()
